#!/usr/bin/env python3
"""
Tesla OTA Update Classifier Pipeline v2
Bachelor Thesis - Amine
============================================================
This script:
1. Loads Tesla OTA release notes from Excel
2. Extracts function names from release notes
3. Classifies each function using 3 trained BERT models:
   - Domain (Powertrain, Chassis, Body, HMI, Driving Automation)
   - Function Type (NEW, MODIFIED, REMOVED)
   - Coordination (STANDALONE, COORD_ACTIVE)
4. Exports results to Excel and CSV
"""

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
import torch
from transformers import BertTokenizer, BertForSequenceClassification
import pickle
import re

# ============================================================
# CONFIGURATION
# ============================================================
RELEASE_NOTES_FILE = "20260525_Amine_First_Scrap.xlsx"
DOMAIN_MODEL_PATH = "domain_bert_model.pt"
TYPE_MODEL_PATH = "type_bert_model.pt"
COORD_MODEL_PATH = "coord_bert_model.pt"
DOMAIN_ENCODER_PATH = "domain_label_encoder.pkl"
TYPE_ENCODER_PATH = "type_label_encoder.pkl"
COORD_ENCODER_PATH = "coord_label_encoder.pkl"
OUTPUT_EXCEL = "tesla_ota_classified_v2.xlsx"
OUTPUT_CSV = "tesla_ota_classified_v2.csv"

# ============================================================
# STEP 1: Load models
# ============================================================
def load_models():
    print("Loading BERT models...")
    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    tokenizer = BertTokenizer.from_pretrained("bert-base-uncased")

    with open(DOMAIN_ENCODER_PATH, "rb") as f:
        le_domain = pickle.load(f)
    with open(TYPE_ENCODER_PATH, "rb") as f:
        le_type = pickle.load(f)
    with open(COORD_ENCODER_PATH, "rb") as f:
        le_coord = pickle.load(f)

    domain_model = BertForSequenceClassification.from_pretrained(
        "bert-base-uncased", num_labels=len(le_domain.classes_)
    )
    domain_model.load_state_dict(torch.load(DOMAIN_MODEL_PATH, map_location=device))
    domain_model.to(device).eval()

    type_model = BertForSequenceClassification.from_pretrained(
        "bert-base-uncased", num_labels=len(le_type.classes_)
    )
    type_model.load_state_dict(torch.load(TYPE_MODEL_PATH, map_location=device))
    type_model.to(device).eval()

    coord_model = BertForSequenceClassification.from_pretrained(
        "bert-base-uncased", num_labels=len(le_coord.classes_)
    )
    coord_model.load_state_dict(torch.load(COORD_MODEL_PATH, map_location=device))
    coord_model.to(device).eval()

    print(f"  Domain classes: {list(le_domain.classes_)}")
    print(f"  Type classes: {list(le_type.classes_)}")
    print(f"  Coord classes: {list(le_coord.classes_)}")
    print(f"  Device: {device}")
    return domain_model, type_model, coord_model, le_domain, le_type, le_coord, tokenizer, device

# ============================================================
# STEP 2: Predict function
# ============================================================
def predict(model, le, text, tokenizer, device):
    encoding = tokenizer(
        text, max_length=64, padding="max_length",
        truncation=True, return_tensors="pt"
    )
    input_ids = encoding["input_ids"].to(device)
    attention_mask = encoding["attention_mask"].to(device)
    with torch.no_grad():
        outputs = model(input_ids, attention_mask=attention_mask)
        pred = torch.argmax(outputs.logits, dim=1).item()
    return le.inverse_transform([pred])[0]

# ============================================================
# STEP 3: Extract function names from release notes
# ============================================================
def extract_functions(release_note):
    if not release_note:
        return []
    pattern = r"\[([^\]]+)\]"
    functions = re.findall(pattern, str(release_note))
    seen = set()
    unique = []
    for f in functions:
        if f not in seen:
            seen.add(f)
            unique.append(f)
    return unique

# ============================================================
# STEP 4: Load release notes
# ============================================================
def load_release_notes():
    print("\nLoading release notes...")
    wb = load_workbook(RELEASE_NOTES_FILE, read_only=True)
    ws = wb.active
    data = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        update_id = row[0]
        update_name = row[1]
        release_date = row[2]
        release_note = row[7]
        if update_id and release_note:
            data.append({
                "update_id": update_id,
                "update_name": str(update_name),
                "release_date": str(release_date) if release_date else "",
                "release_note": str(release_note)
            })
    print(f"  Loaded {len(data)} updates with release notes")
    return data

# ============================================================
# STEP 5: Classify all functions
# ============================================================
def classify_all(updates, domain_model, type_model, coord_model, le_domain, le_type, le_coord, tokenizer, device):
    print("\nClassifying functions...")
    results = []
    function_counter = {}

    for update in updates:
        update_id = update["update_id"]
        functions = extract_functions(update["release_note"])

        if not functions:
            continue

        function_counter[update_id] = 0
        for func_name in functions:
            function_counter[update_id] += 1
            f_num = function_counter[update_id]
            function_id = f"{update_id}_F{f_num:03d}"

            domain = predict(domain_model, le_domain, func_name, tokenizer, device)
            ftype = predict(type_model, le_type, func_name, tokenizer, device)
            coord = predict(coord_model, le_coord, func_name, tokenizer, device)

            results.append({
                "function_id": function_id,
                "update_id": update_id,
                "update_name": update["update_name"],
                "release_date": update["release_date"],
                "function_name": func_name,
                "domain_primary": domain,
                "function_type": ftype,
                "standalone_or_coord": coord,
                "classified_by": "BERT_AUTO"
            })
            print(f"  {function_id}: {func_name} -> {domain} | {ftype} | {coord}")

    print(f"\nTotal functions classified: {len(results)}")
    return results

# ============================================================
# STEP 6: Export to Excel and CSV
# ============================================================
def export_results(results):
    print(f"\nExporting results...")
    df = pd.DataFrame(results)

    domain_colors = {
        "Driving Automation": "D5E8D4",
        "HMI/Multimedia/Telematics": "DAE8FC",
        "Powertrain": "FFE6CC",
        "Body": "E1D5E7",
        "Chassis": "FFF2CC",
    }
    type_colors = {
        "NEW": "C6EFCE",
        "MODIFIED": "FFEB9C",
        "REMOVED": "FFC7CE",
    }
    coord_colors = {
        "STANDALONE": "F2F2F2",
        "COORD_ACTIVE": "FFD700",
    }

    with pd.ExcelWriter(OUTPUT_EXCEL, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="ML_CLASSIFICATIONS")
        ws = writer.sheets["ML_CLASSIFICATIONS"]

        summary = df.groupby(["domain_primary", "function_type", "standalone_or_coord"]).size().reset_index(name="count")
        summary.to_excel(writer, index=False, sheet_name="SUMMARY")

        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            domain_val = ws.cell(row=row_idx, column=6).value
            type_val = ws.cell(row=row_idx, column=7).value
            coord_val = ws.cell(row=row_idx, column=8).value
            if domain_val in domain_colors:
                ws.cell(row=row_idx, column=6).fill = PatternFill("solid", start_color=domain_colors[domain_val])
            if type_val in type_colors:
                ws.cell(row=row_idx, column=7).fill = PatternFill("solid", start_color=type_colors[type_val])
            if coord_val in coord_colors:
                ws.cell(row=row_idx, column=8).fill = PatternFill("solid", start_color=coord_colors[coord_val])

    df.to_csv(OUTPUT_CSV, index=False)
    print(f"✅ Excel: {OUTPUT_EXCEL}")
    print(f"✅ CSV: {OUTPUT_CSV}")
    return df

# ============================================================
# MAIN
# ============================================================
if __name__ == "__main__":
    print("=" * 60)
    print("Tesla OTA Update Classifier Pipeline v2")
    print("=" * 60)

    domain_model, type_model, coord_model, le_domain, le_type, le_coord, tokenizer, device = load_models()
    updates = load_release_notes()
    results = classify_all(updates, domain_model, type_model, coord_model, le_domain, le_type, le_coord, tokenizer, device)
    df = export_results(results)

    print("\n" + "=" * 60)
    print("CLASSIFICATION SUMMARY")
    print("=" * 60)
    print(f"\nDomain distribution:\n{df['domain_primary'].value_counts()}")
    print(f"\nFunction type distribution:\n{df['function_type'].value_counts()}")
    print(f"\nCoordination distribution:\n{df['standalone_or_coord'].value_counts()}")
    print("\n✅ Pipeline complete!")
